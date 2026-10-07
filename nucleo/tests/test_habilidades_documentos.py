import json

import pytest
from docx import Document

from azul.adapters.documentos_windows import DocumentosWindows
from azul.adapters.habilidades_archivos import cargar_habilidades
from azul.config import REPO_ROOT
from azul.core.herramientas_documentos import herramienta_habilidades, herramientas_documentos
from azul.core.persona import build_system_prompt
from azul.core.ports import DocumentosError, Fact, Habilidad

pytestmark = pytest.mark.anyio

REDACCION = Habilidad("redaccion", "Redactar textos.", "# Redacción\nNo inventes datos.")


# --- Habilidades ---


def test_skills_are_read_from_skill_md_files(tmp_path):
    carpeta = tmp_path / "habilidades" / "juridica"
    carpeta.mkdir(parents=True)
    (carpeta / "SKILL.md").write_text(
        "---\nname: juridica\ndescription: Consultas de derecho colombiano.\n---\n\n# Jurídica\n",
        encoding="utf-8",
    )
    (tmp_path / "habilidades" / "rota").mkdir()
    (tmp_path / "habilidades" / "rota" / "SKILL.md").write_text("sin encabezado", encoding="utf-8")

    [habilidad] = cargar_habilidades(tmp_path / "habilidades")

    assert habilidad == Habilidad("juridica", "Consultas de derecho colombiano.", "# Jurídica")


def test_azul_ships_with_its_skills():
    nombres = {h.nombre for h in cargar_habilidades(REPO_ROOT / "habilidades")}

    assert nombres == {
        "redaccion",
        "juridica",
        "lectura",
        "excel",
        "diseno-interfaces",
        "programacion",
    }


def test_system_prompt_lists_skills_but_not_their_instructions():
    prompt = build_system_prompt([Fact("Se llama Camilo.")], [REDACCION])

    assert "- redaccion: Redactar textos." in prompt
    assert "No inventes datos." not in prompt
    assert prompt.index("redaccion") < prompt.index("Se llama Camilo.")


async def test_use_skill_tool_returns_the_instructions():
    herramienta = herramienta_habilidades([REDACCION])

    assert await herramienta.handler({"nombre": "redaccion"}) == REDACCION.instrucciones
    with pytest.raises(ValueError):
        await herramienta.handler({"nombre": "otra"})
    assert herramienta_habilidades([]) is None


# --- Documentos (adaptador) ---


def make_documentos(tmp_path):
    raiz = tmp_path / "pc"
    (raiz / "Trabajo").mkdir(parents=True)
    (raiz / "AppData").mkdir()
    return DocumentosWindows([raiz], tmp_path / "OneDrive" / "Azul" / "Documentos"), raiz


def word(ruta, *parrafos, encabezado=None):
    documento = Document()
    if encabezado:
        documento.sections[0].header.paragraphs[0].text = encabezado
    for parrafo in parrafos:
        documento.add_paragraph(parrafo)
    documento.save(str(ruta))
    return ruta


async def test_search_by_words_of_the_name_ignoring_accents(tmp_path):
    documentos, raiz = make_documentos(tmp_path)
    word(raiz / "Trabajo" / "Carta Terminación Pedro.docx", "x")
    (raiz / "Trabajo" / "notas.txt").write_text("x", encoding="utf-8")
    (raiz / "AppData" / "carta terminacion pedro.docx").write_text("x", encoding="utf-8")

    encontrados = await documentos.buscar("carta terminacion", ["docx"])

    assert [r["ruta"] for r in encontrados] == [
        str(raiz / "Trabajo" / "Carta Terminación Pedro.docx")
    ]


async def test_reads_word_text_and_tables(tmp_path):
    documentos, raiz = make_documentos(tmp_path)
    ruta = word(raiz / "acta.docx", "Acta de reunión", "Asistentes: Camilo")
    documento = Document(str(ruta))
    tabla = documento.add_table(rows=1, cols=2)
    tabla.cell(0, 0).text, tabla.cell(0, 1).text = "IPS", "Cali"
    documento.save(str(ruta))

    texto = await documentos.leer(str(ruta))

    assert "Acta de reunión" in texto and "IPS | Cali" in texto


async def test_reads_pdf_text(tmp_path):
    documentos, raiz = make_documentos(tmp_path)
    ruta = raiz / "contrato.pdf"
    ruta.write_bytes(_pdf_con_texto("Contrato de prueba"))

    assert "Contrato de prueba" in await documentos.leer(str(ruta))


async def test_refuses_to_read_files_that_look_like_secrets(tmp_path):
    documentos, raiz = make_documentos(tmp_path)
    for nombre in ("KEY Anthropic.txt", ".env", "mis contraseñas.txt"):
        (raiz / nombre).write_text("secreto", encoding="utf-8")
        with pytest.raises(DocumentosError, match="claves"):
            await documentos.leer(str(raiz / nombre))


async def test_creates_word_in_onedrive_and_never_overwrites(tmp_path):
    documentos, _ = make_documentos(tmp_path)
    contenido = "# Informe\n\nPrimer párrafo.\n\n- uno\n- dos"

    primera = await documentos.crear_word("Informe: septiembre", contenido, None)
    segunda = await documentos.crear_word("Informe: septiembre", contenido, None)

    assert primera.endswith("Informe septiembre.docx")
    assert segunda.endswith("Informe septiembre (2).docx")
    parrafos = Document(primera).paragraphs
    assert [(p.text, p.style.name) for p in parrafos] == [
        ("Informe", "Heading 1"),
        ("Primer párrafo.", "Normal"),
        ("uno", "List Bullet"),
        ("dos", "List Bullet"),
    ]


async def test_word_from_a_model_keeps_its_letterhead_and_replaces_the_body(tmp_path):
    documentos, raiz = make_documentos(tmp_path)
    modelo = word(raiz / "formato.docx", "Texto viejo", encabezado="CONFIANZA IPS")

    ruta = await documentos.crear_word("Carta", "Texto nuevo.", str(modelo))

    nuevo = Document(ruta)
    assert nuevo.sections[0].header.paragraphs[0].text == "CONFIANZA IPS"
    assert [p.text for p in nuevo.paragraphs] == ["Texto nuevo."]
    assert [p.text for p in Document(str(modelo)).paragraphs] == ["Texto viejo"]


# --- Herramientas de documentos (núcleo) ---


class FakeDocumentos:
    async def buscar(self, texto, extensiones):
        return [{"ruta": f"C:\\{texto}.docx"}] if texto != "nada" else []

    async def leer(self, ruta):
        raise DocumentosError("No encuentro ese archivo.")

    async def pdf_a_word(self, ruta):
        return ruta.replace(".pdf", ".docx")

    async def crear_word(self, titulo, contenido, modelo):
        return f"C:\\OneDrive\\Azul\\Documentos\\{titulo}.docx"

    async def crear_excel(self, titulo, hojas):
        return f"C:\\OneDrive\\Azul\\Documentos\\{titulo}.xlsx"

    async def guardar_archivo(self, nombre, extension, contenido):
        return f"C:\\OneDrive\\Azul\\Documentos\\{nombre}.{extension}"


async def test_document_tools_report_results_and_errors():
    anotado = []
    buscar, leer, convertir, crear, crear_excel, guardar = herramientas_documentos(
        FakeDocumentos(), anotado.append
    )

    assert json.loads(await buscar.handler({"texto": "carta", "tipos": ["docx", "exe"]}))
    assert await buscar.handler({"texto": "nada"}) == "No encontré archivos con ese nombre."
    with pytest.raises(ValueError, match="No encuentro"):
        await leer.handler({"ruta": "C:\\x.docx"})
    assert (await convertir.handler({"ruta": "C:\\a.pdf"})).endswith("C:\\a.docx")
    assert "Carta.docx" in await crear.handler({"titulo": "Carta", "contenido": "Hola."})
    hojas = [{"nombre": "Datos", "filas": [["A"], [1]]}]
    assert "Cuadro.xlsx" in await crear_excel.handler({"titulo": "Cuadro", "hojas": hojas})
    assert "tablero.html" in await guardar.handler(
        {"nombre": "tablero", "extension": "html", "contenido": "<h1>Hola</h1>"}
    )
    with pytest.raises(ValueError):
        await crear_excel.handler({"titulo": "Vacío", "hojas": []})
    assert anotado == ["archivos del PC"] * 6


def _pdf_con_texto(texto: str) -> bytes:
    """Un PDF mínimo de una página con una línea de texto."""
    flujo = f"BT /F1 12 Tf 72 720 Td ({texto}) Tj ET".encode()
    objetos = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(flujo)).encode() + b" >>\nstream\n" + flujo + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    salida = bytearray(b"%PDF-1.4\n")
    posiciones = []
    for numero, cuerpo in enumerate(objetos, start=1):
        posiciones.append(len(salida))
        salida += f"{numero} 0 obj\n".encode() + cuerpo + b"\nendobj\n"
    inicio = len(salida)
    salida += f"xref\n0 {len(objetos) + 1}\n0000000000 65535 f \n".encode()
    for posicion in posiciones:
        salida += f"{posicion:010d} 00000 n \n".encode()
    salida += (
        f"trailer << /Size {len(objetos) + 1} /Root 1 0 R >>\nstartxref\n{inicio}\n%%EOF".encode()
    )
    return bytes(salida)


async def test_creates_excel_with_header_numbers_and_formulas(tmp_path):
    from openpyxl import load_workbook

    documentos, _ = make_documentos(tmp_path)
    hojas = [
        {
            "nombre": "Ventas",
            "filas": [["Ciudad", "Valor"], ["Cali", 1500000], ["Total", "=SUM(B2:B2)"]],
        },
        {"nombre": "Ventas", "filas": [["Otra"]]},
    ]

    ruta = await documentos.crear_excel("Ventas octubre", hojas)

    libro = load_workbook(ruta)
    assert libro.sheetnames == ["Ventas", "Ventas 2"]
    hoja = libro["Ventas"]
    assert hoja["B2"].value == 1500000
    assert hoja["B3"].value == "=SUM(B2:B2)"
    assert hoja["A1"].font.bold
    assert hoja.freeze_panes == "A2"


async def test_reads_excel_sheets_as_rows(tmp_path):
    documentos, _ = make_documentos(tmp_path)
    filas = [["IPS", "Ciudad"], ["Sana", "Cali"]]
    ruta = await documentos.crear_excel("Datos", [{"nombre": "IPS", "filas": filas}])

    texto = await documentos.leer(ruta)

    assert "Hoja «IPS»" in texto
    assert "Sana | Cali" in texto


async def test_saves_pages_and_code_but_not_other_types(tmp_path):
    documentos, _ = make_documentos(tmp_path)

    ruta = await documentos.guardar_archivo("Tablero", "html", "<!doctype html><h1>Hola</h1>")

    assert ruta.endswith("Tablero.html")
    assert "Hola" in await documentos.leer(ruta)
    with pytest.raises(DocumentosError):
        await documentos.guardar_archivo("virus", "exe", "x")
