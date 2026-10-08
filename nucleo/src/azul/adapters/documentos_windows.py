"""Los archivos del usuario en el portátil (ADR 0032, 0034).

- Buscar: por palabras del nombre, con un tiempo máximo, saltando carpetas del
  sistema. Por decisión del usuario puede buscar en todo el equipo.
- Leer: Word (python-docx), PDF (pypdf), Excel (openpyxl), texto y código.
- PDF a Word: con el Microsoft Word del usuario (PowerShell + COM), que es el
  que mejor conserva el formato.
- Crear Word, Excel y archivos de texto o código (HTML, Python…) en
  OneDrive/Azul/Documentos; un Word con modelo conserva su membrete.

Nunca borra ni sobrescribe. No lee archivos que parezcan secretos (claves, .env).
"""

import asyncio
import logging
import os
import re
import subprocess
import time
import unicodedata
from datetime import datetime
from pathlib import Path

from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pypdf import PdfReader

from azul.adapters.archivos_basicos import parece_secreto
from azul.core.ports import DocumentosError

log = logging.getLogger(__name__)

MAX_RESULTADOS = 15
SEGUNDOS_BUSCANDO = 10.0
MAX_CARACTERES = 30_000
MAX_PAGINAS_PDF = 40
SEGUNDOS_CONVIRTIENDO = 180

# Texto plano: notas, datos y código fuente.
EXTENSIONES_TEXTO = {
    ".txt",
    ".md",
    ".csv",
    ".json",
    ".xml",
    ".yaml",
    ".yml",
    ".ini",
    ".log",
    ".html",
    ".htm",
    ".css",
    ".js",
    ".ts",
    ".tsx",
    ".jsx",
    ".py",
    ".sql",
    ".bat",
    ".ps1",
    ".sh",
    ".java",
    ".cs",
    ".php",
    ".go",
    ".rb",
    ".vue",
}
EXTENSIONES_LEGIBLES = {".docx", ".pdf", ".xlsx", ".xlsm", *EXTENSIONES_TEXTO}
# Lo que Azul puede escribir como archivo de texto (páginas, código, datos).
EXTENSIONES_ESCRIBIBLES = {
    "html",
    "css",
    "js",
    "ts",
    "py",
    "md",
    "txt",
    "json",
    "csv",
    "sql",
    "xml",
    "yaml",
}
MAX_FILAS_EXCEL = 200
# Carpetas que no tiene sentido recorrer (sistema, programas, cachés).
_SALTAR = {
    "windows",
    "program files",
    "program files (x86)",
    "programdata",
    "appdata",
    "$recycle.bin",
    "system volume information",
    "node_modules",
    ".git",
    ".venv",
    "__pycache__",
    "recovery",
    "perflogs",
}
# Nombres que parecen secretos: Azul no los lee aunque se lo pidan.


class DocumentosWindows:
    def __init__(self, raices: list[Path], salida: Path) -> None:
        # Las raíces se recorren en orden: primero las más probables.
        self._raices = [r for r in raices if r.is_dir()]
        self._salida = salida

    async def buscar(self, texto: str, extensiones: list[str]) -> list[dict[str, object]]:
        return await asyncio.to_thread(self._buscar, texto, extensiones)

    async def leer(self, ruta: str) -> str:
        return await asyncio.to_thread(self._leer, Path(ruta))

    async def pdf_a_word(self, ruta: str) -> str:
        return await asyncio.to_thread(self._pdf_a_word, Path(ruta))

    async def crear_excel(self, titulo: str, hojas: list[dict[str, object]]) -> str:
        return await asyncio.to_thread(self._crear_excel, titulo, hojas)

    async def guardar_archivo(self, nombre: str, extension: str, contenido: str) -> str:
        return await asyncio.to_thread(self._guardar_archivo, nombre, extension, contenido)

    async def crear_word(self, titulo: str, contenido: str, modelo: str | None) -> str:
        return await asyncio.to_thread(
            self._crear_word, titulo, contenido, Path(modelo) if modelo else None
        )

    # --- Buscar ---

    def _buscar(self, texto: str, extensiones: list[str]) -> list[dict[str, object]]:
        palabras = _normalizar(texto).split()
        sufijos = tuple(f".{e.lower()}" for e in extensiones) if extensiones else None
        limite = time.monotonic() + SEGUNDOS_BUSCANDO
        vistos: set[str] = set()
        encontrados: list[Path] = []
        for raiz in self._raices:
            for carpeta, subcarpetas, archivos in os.walk(raiz):
                subcarpetas[:] = [
                    s for s in subcarpetas if s.lower() not in _SALTAR and not s.startswith(".")
                ]
                for nombre in archivos:
                    if sufijos and not nombre.lower().endswith(sufijos):
                        continue
                    plano = _normalizar(nombre)
                    if all(p in plano for p in palabras):
                        ruta = Path(carpeta, nombre)
                        clave = str(ruta).lower()
                        if clave not in vistos:
                            vistos.add(clave)
                            encontrados.append(ruta)
                if len(encontrados) >= MAX_RESULTADOS * 3 or time.monotonic() > limite:
                    break
        encontrados.sort(key=_fecha, reverse=True)
        return [_ficha(r) for r in encontrados[:MAX_RESULTADOS]]

    # --- Leer ---

    def _leer(self, ruta: Path) -> str:
        if parece_secreto(ruta.name):
            raise DocumentosError("Ese archivo parece guardar claves; por seguridad no lo leo.")
        if not ruta.is_file():
            raise DocumentosError("No encuentro ese archivo.")
        sufijo = ruta.suffix.lower()
        if sufijo not in EXTENSIONES_LEGIBLES:
            raise DocumentosError(f"Aún no sé leer archivos {sufijo or 'sin extensión'}.")
        try:
            if sufijo == ".docx":
                texto = _texto_de_word(ruta)
            elif sufijo == ".pdf":
                texto = _texto_de_pdf(ruta)
            elif sufijo in (".xlsx", ".xlsm"):
                texto = _texto_de_excel(ruta)
            else:
                texto = ruta.read_text(encoding="utf-8", errors="replace")
        except DocumentosError:
            raise
        except Exception as error:
            log.warning("No se pudo leer un %s: %s", sufijo, type(error).__name__)
            raise DocumentosError(
                "No pude abrir ese archivo; puede estar dañado o protegido."
            ) from error
        if not texto.strip():
            return "El archivo no tiene texto legible (puede ser un escaneo o una imagen)."
        if len(texto) > MAX_CARACTERES:
            return texto[:MAX_CARACTERES] + "\n…(el documento sigue; esto es lo primero)"
        return texto

    # --- PDF a Word ---

    def _pdf_a_word(self, ruta: Path) -> str:
        if ruta.suffix.lower() != ".pdf" or not ruta.is_file():
            raise DocumentosError("No encuentro ese PDF.")
        destino = _libre(self._carpeta_salida() / f"{ruta.stem}.docx")
        # Word abre el PDF, lo convierte a documento editable y lo guarda como .docx.
        guion = (
            "$ErrorActionPreference='Stop';"
            "$w=New-Object -ComObject Word.Application;$w.Visible=$false;$w.DisplayAlerts=0;"
            "try{$d=$w.Documents.Open($env:AZUL_PDF,$false,$true);"
            "$d.SaveAs2($env:AZUL_DOCX,16);$d.Close($false)}finally{$w.Quit()}"
        )
        try:
            subprocess.run(
                ["powershell", "-NoProfile", "-NonInteractive", "-Command", guion],
                env={**os.environ, "AZUL_PDF": str(ruta), "AZUL_DOCX": str(destino)},
                capture_output=True,
                timeout=SEGUNDOS_CONVIRTIENDO,
                check=True,
                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
            )
        except (subprocess.SubprocessError, OSError) as error:
            log.warning("Word no pudo convertir el PDF: %s", type(error).__name__)
            raise DocumentosError(
                "Word no pudo convertir ese PDF. ¿Está instalado y sin ventanas abiertas esperando?"
            ) from error
        if not destino.is_file():
            raise DocumentosError("Word terminó pero no dejó el archivo convertido.")
        return str(destino)

    # --- Crear Word ---

    def _crear_word(self, titulo: str, contenido: str, modelo: Path | None) -> str:
        if modelo is not None:
            if modelo.suffix.lower() != ".docx" or not modelo.is_file():
                raise DocumentosError("El modelo tiene que ser un Word (.docx) que exista.")
            documento = Document(str(modelo))
            _vaciar_cuerpo(documento)
        else:
            documento = Document()
        estilos = {estilo.name for estilo in documento.styles}
        for bloque in re.split(r"\n\s*\n", contenido.strip()):
            for parrafo in _parrafos(bloque, estilos):
                texto, estilo = parrafo
                documento.add_paragraph(texto, style=estilo)
        destino = _libre(self._carpeta_salida() / f"{_nombre_seguro(titulo)}.docx")
        documento.save(str(destino))
        return str(destino)

    # --- Crear Excel ---

    def _crear_excel(self, titulo: str, hojas: list[dict[str, object]]) -> str:
        if not hojas:
            raise DocumentosError("El Excel necesita al menos una hoja con filas.")
        libro = Workbook()
        libro.remove(libro.active)
        usados: set[str] = set()
        for numero, hoja in enumerate(hojas, start=1):
            nombre = _nombre_de_hoja(str(hoja.get("nombre") or f"Hoja{numero}"), usados)
            filas = hoja.get("filas") or []
            if not isinstance(filas, list):
                raise DocumentosError("Las filas de una hoja deben ser una lista.")
            _llenar_hoja(libro.create_sheet(nombre), filas)
        destino = _libre(self._carpeta_salida() / f"{_nombre_seguro(titulo)}.xlsx")
        libro.save(str(destino))
        return str(destino)

    # --- Guardar texto o código ---

    def _guardar_archivo(self, nombre: str, extension: str, contenido: str) -> str:
        extension = extension.lower().lstrip(".")
        if extension not in EXTENSIONES_ESCRIBIBLES:
            raise DocumentosError(f"No guardo archivos .{extension}.")
        destino = _libre(self._carpeta_salida() / f"{_nombre_seguro(nombre)}.{extension}")
        destino.write_text(contenido, encoding="utf-8")
        return str(destino)

    def _carpeta_salida(self) -> Path:
        self._salida.mkdir(parents=True, exist_ok=True)
        return self._salida


def raices_por_defecto() -> list[Path]:
    """Primero lo más probable (escritorio, documentos, descargas, OneDrive), luego todo."""
    casa = Path.home()
    probables = [
        Path(r"C:\SANDRA\Desktop"),
        casa / "Desktop",
        casa / "Documents",
        casa / "Downloads",
        Path(os.environ.get("ONEDRIVE", casa / "OneDrive")),
        casa,
    ]
    unidades = [Path(f"{letra}:\\") for letra in "CDEFG" if Path(f"{letra}:\\").is_dir()]
    vistas: list[Path] = []
    for ruta in [*probables, *unidades]:
        if ruta not in vistas:
            vistas.append(ruta)
    return vistas


def _parrafos(bloque: str, estilos: set[str]) -> list[tuple[str, str | None]]:
    lineas = [linea.rstrip() for linea in bloque.splitlines() if linea.strip()]
    if lineas and all(linea.lstrip().startswith("- ") for linea in lineas):
        estilo = "List Bullet" if "List Bullet" in estilos else None
        return [(linea.lstrip()[2:].strip(), estilo) for linea in lineas]
    texto = "\n".join(lineas)
    for marca, nombre in (("## ", "Heading 2"), ("# ", "Heading 1")):
        if texto.startswith(marca):
            return [(texto[len(marca) :].strip(), nombre if nombre in estilos else None)]
    return [(texto, None)]


def _vaciar_cuerpo(documento) -> None:  # type: ignore[no-untyped-def]
    """Quita párrafos y tablas del cuerpo; quedan encabezado, pie, márgenes y estilos."""
    cuerpo = documento.element.body
    for hijo in list(cuerpo):
        if hijo.tag.endswith("}sectPr"):
            continue
        cuerpo.remove(hijo)


def _texto_de_word(ruta: Path) -> str:
    documento = Document(str(ruta))
    partes = [p.text for p in documento.paragraphs]
    for tabla in documento.tables:
        for fila in tabla.rows:
            partes.append(" | ".join(celda.text.strip() for celda in fila.cells))
    return "\n".join(partes)


def _texto_de_excel(ruta: Path) -> str:
    """Cada hoja con sus primeras filas, celdas separadas por " | " (valores, no fórmulas)."""
    libro = load_workbook(str(ruta), read_only=True, data_only=True)
    partes = []
    try:
        for hoja in libro.worksheets:
            partes.append(
                f"## Hoja «{hoja.title}» ({hoja.max_row} filas x {hoja.max_column} columnas)"
            )
            for numero, fila in enumerate(hoja.iter_rows(values_only=True), start=1):
                if numero > MAX_FILAS_EXCEL:
                    partes.append(f"…(solo las primeras {MAX_FILAS_EXCEL} filas)")
                    break
                if any(celda is not None for celda in fila):
                    partes.append(" | ".join("" if c is None else str(c) for c in fila))
    finally:
        libro.close()
    return "\n".join(partes)


def _llenar_hoja(hoja, filas: list[object]) -> None:  # type: ignore[no-untyped-def]
    """Primera fila como encabezado (negrita, fondo, fija al desplazarse); anchos ajustados."""
    anchos: dict[int, int] = {}
    for numero, fila in enumerate(filas, start=1):
        celdas = fila if isinstance(fila, list) else [fila]
        for columna, valor in enumerate(celdas, start=1):
            celda = hoja.cell(row=numero, column=columna, value=_valor(valor))
            if numero == 1:
                celda.font = Font(bold=True, color="FFFFFF")
                celda.fill = PatternFill("solid", fgColor="1F3A66")
                celda.alignment = Alignment(vertical="center", wrap_text=True)
            anchos[columna] = max(anchos.get(columna, 0), len(str(valor or "")))
    for columna, ancho in anchos.items():
        hoja.column_dimensions[get_column_letter(columna)].width = min(60, max(10, ancho + 2))
    if len(filas) > 1:
        hoja.freeze_panes = "A2"


def _valor(valor: object) -> object:
    # Las fórmulas ("=SUMA…") y los números se guardan tal cual; lo demás, como texto.
    if valor is None or isinstance(valor, (int, float, bool)):
        return valor
    return str(valor)


def _nombre_de_hoja(nombre: str, usados: set[str]) -> str:
    limpio = re.sub(r"[\[\]:*?/\\]", "", nombre).strip()[:31] or "Hoja"
    candidato, numero = limpio, 2
    while candidato.lower() in usados:
        candidato = f"{limpio[:28]} {numero}"
        numero += 1
    usados.add(candidato.lower())
    return candidato


def _texto_de_pdf(ruta: Path) -> str:
    lector = PdfReader(str(ruta))
    if lector.is_encrypted:
        raise DocumentosError("Ese PDF tiene contraseña; no puedo abrirlo.")
    paginas = [p.extract_text() or "" for p in lector.pages[:MAX_PAGINAS_PDF]]
    texto = "\n\n".join(paginas)
    if len(lector.pages) > MAX_PAGINAS_PDF:
        texto += f"\n…(solo las primeras {MAX_PAGINAS_PDF} de {len(lector.pages)} páginas)"
    return texto


def _normalizar(texto: str) -> str:
    plano = unicodedata.normalize("NFD", texto.lower())
    return "".join(c for c in plano if unicodedata.category(c) != "Mn")


def _fecha(ruta: Path) -> float:
    try:
        return ruta.stat().st_mtime
    except OSError:
        return 0.0


def _ficha(ruta: Path) -> dict[str, object]:
    try:
        estado = ruta.stat()
        return {
            "ruta": str(ruta),
            "kb": round(estado.st_size / 1024),
            "modificado": datetime.fromtimestamp(estado.st_mtime).strftime("%Y-%m-%d %H:%M"),
        }
    except OSError:
        return {"ruta": str(ruta)}


def _nombre_seguro(titulo: str) -> str:
    limpio = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", titulo).strip().rstrip(".")
    return limpio[:80] or "Documento de Azul"


def _libre(destino: Path) -> Path:
    """Nunca sobrescribe: si el nombre existe, agrega (2), (3)…"""
    if not destino.exists():
        return destino
    for numero in range(2, 1000):
        otro = destino.with_name(f"{destino.stem} ({numero}){destino.suffix}")
        if not otro.exists():
            return otro
    raise DocumentosError("Hay demasiadas versiones con ese nombre; usa otro título.")
